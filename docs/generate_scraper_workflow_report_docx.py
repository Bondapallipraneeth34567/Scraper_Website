"""Renders docs/mandi_scraper_workflow.docx: the whole AGMARKNET scraper
(mandi_pipeline, Approach 1) end to end -- triggers, sweep, CAPTCHA,
fetch, storage, CSV/blob, monitoring, verification and the QA dashboard --
plus the problems met while building/running it and how each was fixed,
and a full workflow diagram at the end.

Run: python docs/generate_scraper_workflow_report_docx.py
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_PATH = Path(__file__).parent / "mandi_scraper_workflow.docx"
DIAGRAM_PATH = Path(__file__).parent / "mandi_scraper_workflow_diagram.png"

CODE_FONT = "Consolas"
BODY_FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x5C, 0x3D)  # dark green, agri-themed
DARK_TEXT = RGBColor(0x22, 0x22, 0x22)


# ------------------------------------------------------------------ helpers
def set_shading(pr, fill):
    pr.append(pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill,
    }))


def add_code_block(doc, text):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        set_shading(p._p.get_or_add_pPr(), "F2F2F2")
        run = p.add_run(line if line.strip() else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(10.5)
    return p


def add_bullet(doc, label, rest):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if label:
        r1 = p.add_run(label)
        r1.bold = True
        r1.font.name = BODY_FONT
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = ACCENT
    r2 = p.add_run(rest)
    r2.font.name = BODY_FONT
    r2.font.size = Pt(10.5)


def add_step(doc, n, title, text):
    # explicit numbers: Word's "List Number" style keeps counting across
    # sections, so each section's steps would not restart at 1
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r1 = p.add_run(f"{n}.  {title}. ")
    r1.bold = True
    r1.font.color.rgb = ACCENT
    p.add_run(text)


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
        run.font.name = BODY_FONT
    return h


def add_table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_shading(cell._tc.get_or_add_tcPr(), "1F5C3D")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ------------------------------------------------------------------ diagram
def draw_diagram(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 15))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 150)
    ax.axis("off")
    green, light, blue, amber, grey = "#1F5C3D", "#E8F1EC", "#E6EEF7", "#FFF4E0", "#F3F3F3"

    def box(x, y, w, h, text, fc=light, bold=False):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.5",
                                    fc=fc, ec=green, lw=1.2))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=8.2,
                fontweight="bold" if bold else "normal")
        return (x, y, w, h)

    def down(a, b, label=""):
        x0 = a[0] + a[2] / 2
        x1 = b[0] + b[2] / 2
        ax.annotate("", xy=(x1, b[1] + b[3]), xytext=(x0, a[1]),
                    arrowprops=dict(arrowstyle="-|>", color=green, lw=1.2))
        if label:
            ax.text((x0 + x1) / 2 + 1, (a[1] + b[1] + b[3]) / 2, label, fontsize=7.4,
                    color="#555555", va="center")

    def side(a, b, label="", back=False, y=None):
        """horizontal arrow from a's right edge to b's left edge (or reverse)."""
        yy = y if y is not None else a[1] + a[3] / 2
        if back:
            start, end = (b[0], yy), (a[0] + a[2], yy)
        else:
            start, end = (a[0] + a[2], yy), (b[0], yy)
        ax.annotate("", xy=end, xytext=start,
                    arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
        if label:
            ax.text((start[0] + end[0]) / 2, yy + 1.3, label, fontsize=7.2,
                    color="#555555", ha="center")

    # triggers
    ax.text(2, 148, "Triggers — systemd timers on the VM (UTC)", fontsize=10.5,
            fontweight="bold", color=green, va="top")
    t1 = box(2, 136, 30, 8, "mandi-pipeline.timer\nevery 4 h at :05", blue)
    t2 = box(35, 136, 30, 8, "catch-up timer, daily 02:30\ntarget + previous 3 days", blue)
    t3 = box(68, 136, 30, 8, "log-ship timer, daily 06:20\nrun.py --ship-logs", blue)

    lock = box(20, 124, 45, 7, "Single-instance lock (state/pipeline.lock)", amber)
    ax.annotate("", xy=(30, 131.5), xytext=(17, 136), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.2))
    ax.annotate("", xy=(50, 131.5), xytext=(50, 136), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.2))
    skip = box(1, 113, 16, 8, "lock held:\nskipped_locked,\nexit", grey)
    ax.annotate("", xy=(9, 121.5), xytext=(20, 127.5),
                arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
    ship = box(72, 112, 26, 8, "Upload rotated logs\nto Azure Blob", grey)
    ax.annotate("", xy=(85, 120.5), xytext=(85, 136),
                arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))

    run = box(20, 113, 45, 7, "pipeline_runs row = running\n(zombie 'running' rows -> aborted)")
    tax = box(20, 102, 45, 7, "Taxonomy: GET /filters (no CAPTCHA)\ncached, ~605 commodities")
    dorm = box(20, 91, 45, 7, "Drop dormant commodities\n(always empty 10 d; re-check every 7 d)")
    dates = box(20, 80, 45, 7, "Dates: today (IST), newest first\n+ catch-up days; resume from progress")
    down(lock, run, "acquired")
    down(run, tax)
    down(tax, dorm)
    down(dorm, dates)

    # per-commodity loop frame
    ax.add_patch(FancyBboxPatch((1, 21), 98, 55, boxstyle="round,pad=0.4",
                                fc="#FAFAFA", ec="#999999", lw=1, linestyle="--"))
    ax.text(3, 75, "For each commodity (up to 3 passes)", fontsize=9.5,
            fontweight="bold", color=green, va="top")
    c1 = box(20, 64, 45, 6, "POST /captcha/generator  (image + key)")
    c2 = box(20, 55, 45, 6, "Solve CAPTCHA  (easyocr / trained CNN / LLM)")
    c3 = box(20, 46, 45, 6, "POST /report  (one commodity, one date, limit 100000)")
    c4 = box(20, 37, 45, 6, "Pages via report_ticket -> parse, hash, de-dupe")
    c5 = box(20, 28, 45, 6, "Sync market_prices_v3: insert / update / prune stale")
    c6 = box(20, 22.5, 45, 4, "progress done|empty|failed  ->  CSV  ->  Azure Blob", light)
    down(dates, c1)
    down(c1, c2)
    down(c2, c3, "accepted")
    down(c3, c4, "rows")
    down(c4, c5)
    ax.annotate("", xy=(42.5, 26.5), xytext=(42.5, 28), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.2))

    man = box(72, 55, 26, 7, "4 rejects -> manual\nfallback (90 s)", amber)
    side(c2, man)
    rej = box(72, 64, 26, 6, "server rejected CAPTCHA\n-> new image", amber)
    ax.annotate("", xy=(65, 67), xytext=(72, 67), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
    thr = box(72, 44, 26, 9, "429 / too many / hang:\nbackoff 15 s x 2^n\n+ adaptive spacing", amber)
    side(c3, thr, y=49)
    ax.annotate("", xy=(65, 47), xytext=(72, 47), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
    empty = box(2, 41, 15, 8, "'No data\navailable'\n-> empty", grey)
    ax.annotate("", xy=(17, 45), xytext=(20, 49), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
    dbx = box(72, 27, 26, 8, "Postgres (mandi schema)\nmarket_prices_v3", blue)
    side(c5, dbx, y=31)

    fin = box(20, 10, 45, 7, "finish_run: success | partial | failed + all counters")
    down(c6, fin)
    hb = box(2, 1, 30, 6, "success -> heartbeat ping", grey)
    al = box(35, 1, 30, 6, "partial / failed -> alert webhook", grey)
    down(fin, hb)
    down(fin, al)
    ver = box(72, 10, 26, 9, "verify-csv: site vs DB vs\nCSV / blob ->\nrecord_comparisons", blue)
    dash = box(72, 0.5, 26, 7, "QA dashboard: records,\ncomparison, duplicates", blue)
    ax.annotate("", xy=(85, 19.5), xytext=(85, 27), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))
    ax.annotate("", xy=(85, 8), xytext=(85, 10), arrowprops=dict(arrowstyle="-|>", color=green, lw=1.1))

    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


# ------------------------------------------------------------------ document
def build():
    draw_diagram(DIAGRAM_PATH)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_TEXT
    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    title = doc.add_heading("Mandi Price Scraper — End-to-End Workflow", level=0)
    for run in title.runs:
        run.font.color.rgb = ACCENT
    add_body(doc, "AGMARKNET direct-API pipeline (mandi_pipeline, Approach 1). How it is triggered, how a sweep "
                  "runs, how each commodity is fetched and stored, how the output is checked and shown, the "
                  "problems met along the way and how each was fixed. Updated 01-10-2026.")

    # ---- 1. overview
    add_heading(doc, "1. Overview", 1)
    add_body(doc, "The scraper collects daily mandi (market) prices for every commodity on AGMARKNET and keeps a "
                  "faithful copy of what the site shows. It calls the site's own JSON endpoints over plain HTTP. "
                  "There is no headless browser.")
    add_table(doc, ["Item", "Value"], [
        ["Runs on", "Azure VM sethu-admin1, /opt/mandi_pipeline, as user azureuser"],
        ["Source", "agmarknet.gov.in API: /daily-price-arrival/filters, /captcha/generator, /daily-price-arrival/report"],
        ["Data available from", "07-11-2025 (AGMARKNET 2.0 go-live)"],
        ["Commodities", "~605 in the taxonomy; ~215–280 have prices on a given day"],
        ["Database", "Postgres, schema mandi: market_prices_v3 (prices), pipeline_runs, pipeline_progress, "
                     "agmarknet_taxonomy, record_comparisons, duplicate_audit"],
        ["Files", "One CSV per commodity per day: state/downloads/<date>/<Commodity>.csv, mirrored to Azure Blob "
                  "mandi-prices-csv/agmarknet_api/<date>/"],
        ["CAPTCHA", "One per commodity per date. Production solver: easyocr (CAPTCHA_SOLVER=easyocr)"],
        ["Viewed in", "QA dashboard (backend/ Express API + frontend/ React)"],
    ], widths=[1.6, 5.0])

    # ---- 2. workflow
    add_heading(doc, "2. The workflow, step by step", 1)

    add_heading(doc, "2.1 Triggers", 2)
    add_table(doc, ["systemd timer (UTC)", "Runs", "Purpose"], [
        ["mandi-pipeline.timer — every 4 h at :05", "run.py --once --trigger timer",
         "Sweep today's date. Each tick resumes where the last stopped."],
        ["mandi-pipeline-catchup.timer — daily 02:30", "run.py --once --trigger catchup --catchup-days 3",
         "Re-sweep today plus the previous 3 dates: mandis often enter prices 24–48 h late."],
        ["mandi-pipeline-logship.timer — daily 06:20", "run.py --ship-logs",
         "Upload rotated log archives to blob (agmarknet_api/logs/YYYY-MM/) and free local disk."],
        ["mandi-pipeline-console.service — always on", "operator_console.py",
         "Loopback web page where a person can type a CAPTCHA the solver couldn't read."],
    ], widths=[2.0, 2.2, 2.4])

    add_heading(doc, "2.2 Starting a sweep", 2)
    add_step(doc, 1, "Single-instance lock", "state/pipeline.lock (flock). If another sweep holds it, the run is "
             "recorded as skipped_locked and exits. This is normal when a 4-hour tick overlaps a long sweep.")
    add_step(doc, 2, "Run record", "A mandi.pipeline_runs row is opened with status running. Any older row still "
             "marked running is a sweep that was killed, and is marked aborted.")
    add_step(doc, 3, "Taxonomy", "GET /filters returns every commodity with its id and group id, without a CAPTCHA. "
             "It is cached in mandi.agmarknet_taxonomy; --refresh-taxonomy re-pulls it.")
    add_step(doc, 4, "Dormant commodities", "About 335 of 605 commodities (livestock, wood, exotic flowers) return "
             "'No data available' every day. A commodity empty on every sweep for 10 days is skipped, and re-checked "
             "once every 7 days, saving a CAPTCHA each.")
    add_step(doc, 5, "Dates", "Today in Asia/Kolkata (never before 07-11-2025), plus N earlier dates for catch-up, "
             "newest first. Each date is checkpointed separately.")
    add_step(doc, 6, "Resume", "mandi.pipeline_progress says which commodities are already done or empty for the "
             "date; those are skipped. One finished more than 4 hours ago is swept again to pick up late arrivals.")

    add_heading(doc, "2.3 Fetching one commodity", 2)
    add_step(doc, 1, "Get a CAPTCHA", "POST /captcha/generator returns a base64 PNG and a key valid ~5 minutes. "
             "A 3-second floor separates it from the previous report call.")
    add_step(doc, 2, "Solve it", "The configured solver reads the 6 characters. If the server rejects the answer, "
             "a fresh image is fetched, up to 4 times. After 4 failures the manual fallback shows the image on the "
             "operator console for 90 seconds; if nobody answers, the commodity is retried in the next pass.")
    add_step(doc, 3, "Request the report", "POST /daily-price-arrival/report with the group, commodity, "
             "from_date = to_date = the date, data_type 100006, 'all' sentinels for state/district/market/grade/"
             "variety, limit 100000, and the CAPTCHA answer. A whole commodity normally arrives in one response.")
    add_step(doc, 4, "Interpret the response", "Rows → continue. 'No data available' → recorded as empty (normal). "
             "Body mentions captcha / 'Report filters changed' → new CAPTCHA. HTTP 429, 'too many requests' or a "
             "hung socket → rate limited. Anything else → an optional LLM classifier, otherwise the commodity is "
             "marked failed and the sweep continues.")
    add_step(doc, 5, "Back off when throttled", "Wait 15 s × 2ⁿ (capped at 300 s), up to 4 tries. Each throttle "
             "also widens the gap between report calls by 4 s (up to 30 s); each success narrows it by 1 s. If a "
             "commodity is abandoned for throttling, the sweep pauses 120 s before the next.")
    add_step(doc, 6, "More pages", "If total_pages > 1, the remaining pages are fetched with the response's "
             "report_ticket (30 calls / 5 min, tied to the same filters), without another CAPTCHA.")
    add_step(doc, 7, "Parse", "Tolerant key mapping (state_name, market_name, model_price …), prices and dates "
             "cleaned, a blank price kept blank (never guessed), source_row_hash = SHA-256 of the business fields, "
             "rows de-duplicated on the unique key. The API's declared total is compared with what was fetched and "
             "any shortfall is logged.")
    add_step(doc, 8, "Store", "One transaction per commodity and date: insert new keys, update rows whose hash "
             "changed (first_seen_at kept, scraped_at bumped), leave identical rows untouched, and — only after a "
             "complete fetch — delete rows the site no longer returns. A prune that would remove more than 34% of a "
             "commodity's rows (with at least 15 rows) is skipped as a likely upstream glitch.")
    add_step(doc, 9, "Record progress", "pipeline_progress = done (with row count), empty, or failed (with the "
             "reason).")
    add_step(doc, 10, "Write the CSV", "Rebuilt from the database after the store, in the site's own export "
             "columns and DD-MM-YYYY dates, written atomically, then uploaded to Azure Blob. A blob failure is "
             "logged but does not fail the commodity.")

    add_heading(doc, "2.4 Finishing a sweep", 2)
    add_step(doc, 1, "Passes", "Up to 3 passes over the commodity list; each later pass retries only what failed.")
    add_step(doc, 2, "Run status", "success (nothing failed), partial (some failed), failed (nothing done), or "
             "aborted (an exception). Counters saved: commodities total/done/failed, rows seen/new/changed/"
             "unchanged/missing/pruned, CAPTCHAs solved, throttle events.")
    add_step(doc, 3, "Notify", "On success, HEARTBEAT_URL is pinged (dead-man's switch). On partial or failed, a "
             "JSON summary is posted to ALERT_WEBHOOK_URL.")

    add_heading(doc, "2.5 Checking and viewing the data", 2)
    add_bullet(doc, "Verify: ", "run.py --verify-csv compares, for one date, CSV row counts with the database "
               "(offline). --deep also re-fetches the site and compares row by row and cell by cell; --fix re-scrapes "
               "anything that differs; --blob uses the copy in Azure Blob. Each result goes to "
               "mandi.record_comparisons.")
    add_bullet(doc, "QA dashboard: ", "Data Verification (browse and mark records), Site Comparison (the saved "
               "check results), Duplicates (true duplicates can be merged with an audit log; variety/grade variants "
               "are read-only), Missing Data.")

    add_heading(doc, "2.6 Data model", 2)
    add_table(doc, ["Table", "Holds"], [
        ["market_prices_v3", "Price rows. Unique on (state, district, market, commodity, variety, grade, arrival_date). "
                             "Plus source, source_row_hash, first_seen_at, scraped_at."],
        ["pipeline_runs", "One row per execution, success or failure, with start/finish times and all counters."],
        ["pipeline_progress", "Per date and commodity: done / empty / failed, rows, last error. Drives resume."],
        ["agmarknet_taxonomy", "Cached commodity list (id, name, group)."],
        ["record_comparisons", "Latest site-vs-DB-vs-CSV check per date and commodity."],
        ["duplicate_audit", "Rows deleted from the Duplicates page, kept as JSON."],
    ], widths=[1.7, 4.9])

    # ---- 3. problems
    add_heading(doc, "3. Problems faced and how each was fixed", 1)
    problems = [
        ("One CAPTCHA per commodity",
         "The report_ticket is locked to the exact filter set; changing commodity returns 'Report filters changed. "
         "Please solve captcha again.' That means ~605 CAPTCHAs a day, so a human solver is not practical.",
         "Pluggable solvers: easyocr with image preprocessing (grayscale, upscale, denoise, Otsu, morphology), an "
         "LLM reader, an OCR+LLM hybrid, and a purpose-trained CNN (captcha_model/: 79% whole-string, 96% "
         "per-character on validation). Manual stays as the fallback. Dormant commodities are skipped to cut the "
         "count."),
        ("Rate limiter with no headers",
         "The site's edge limiter sends no Retry-After; after ~4 quick calls the socket just hangs. A polite "
         "60-commodity test solved 52% with 56 throttle events. On 08-09-2026 a full sweep hit 148 backoffs and "
         "abandoned 5 commodities.",
         "Minimum 8 s between report calls, 3 s before each CAPTCHA mint, hung socket treated as throttle, "
         "exponential backoff, multi-pass sweep, and (01-10-2026) adaptive spacing plus a 120 s cooldown after a "
         "throttle give-up."),
        ("Manual fallback blocked the sweep",
         "On 08-09-2026, 9 commodities each waited 30 minutes for an operator who wasn't there (~5 hours lost) and "
         "were still missed.",
         "The fallback after automated solving fails now waits 90 s, then moves on; the next pass retries it."),
        ("Duplicate rows from blank variety/grade",
         "Postgres treats NULL ≠ NULL in a unique constraint, so a same-day re-run of a blank-variety row would "
         "insert a second copy.",
         "variety and grade are NOT NULL DEFAULT ''. The upsert compares a row hash, so unchanged rows aren't "
         "rewritten."),
        ("Late-arriving prices",
         "AGMARKNET fills a day's markets in progressively over 24–48 h, so a sweep early in the day misses many.",
         "A commodity finished more than 4 h ago is swept again on the next tick, and a daily catch-up re-sweeps "
         "the previous 3 dates."),
        ("Stale rows after site changes",
         "AGMARKNET renames markets after the fact (e.g. 'Aaron(F&V) APMC' → 'Aron (F&V) APMC'). The old-name rows "
         "stayed, so the DB held more rows than the site. On 06-09-2026, 102 of 215 commodities were affected.",
         "After a complete fetch, rows the site no longer returns are deleted (mirror prune) with a 34% safety brake. "
         "verify --deep --fix found and repaired the backlog: 06-09-2026 is now 215/215 in sync (61 rows added, "
         "143 stale removed)."),
        ("Partial prices",
         "The site often shows a modal price with min/max blank, or the reverse. Earlier code rejected or guessed "
         "these.",
         "Price columns are nullable and stored exactly as shown; only a row with all three blank is dropped."),
        ("Overlapping and killed runs",
         "A 4-hour timer can fire while a sweep is still going; a killed sweep leaves a 'running' row behind.",
         "flock single-instance lock (overlap = skipped_locked); zombie 'running' rows are marked aborted at the "
         "next start; progress is committed per commodity, so a restart loses nothing."),
        ("Wasted CAPTCHAs on empty commodities",
         "About 335 of 605 commodities return 'No data available' every day, each still costing a CAPTCHA.",
         "Dormant-commodity skip (empty for 10 days → checked once a week)."),
        ("CSV prices rounded",
         "Found 01-10-2026 by the blob-vs-site check: the CSV writer used Python's :g format (6 significant "
         "digits), so 45600.19 was written 45600.2. The database was correct. On 27-09-2026, 17 of 217 "
         "commodities in blob differed from the site for this reason alone.",
         "The writer now prints the stored value exactly. Existing CSVs keep the rounding until rebuilt."),
        ("Misleading duplicate count",
         "The dashboard counted 14,371 'duplicate groups' by ignoring variety and grade; all were separate real "
         "prices (e.g. Grade Range-1/2/3).",
         "True duplicates now include variety and grade (0 found); variants are shown separately, read-only."),
        ("Long checks lost the DB connection",
         "A full-day verify run (~1 h) died when Azure Postgres dropped the connection.",
         "The check pings the connection before each commodity and reconnects if needed."),
        ("Logs filling the VM disk",
         "Rotating logs accumulate on the VM.",
         "Daily log shipping uploads rotated archives to blob and deletes them locally."),
    ]
    add_table(doc, ["Problem", "What happened", "Fix"], [[a, b, c] for a, b, c in problems],
              widths=[1.4, 2.7, 2.6])

    # ---- 4. config
    add_heading(doc, "4. Key settings (.env)", 1)
    add_table(doc, ["Setting", "Default", "Meaning"], [
        ["CAPTCHA_SOLVER", "manual (VM: easyocr)", "manual | easyocr | llm | llm_assisted_ocr | trained | trained_ensemble"],
        ["CAPTCHA_FALLBACK_TIMEOUT_SEC", "90", "Manual wait after the automated solver fails 4 times"],
        ["CAPTCHA_WAIT_TIMEOUT_SEC", "1800", "Wait per CAPTCHA when the solver itself is manual"],
        ["REQUEST_MIN_INTERVAL_SEC", "8", "Minimum gap between report calls"],
        ["BACKOFF_BASE_SEC / BACKOFF_MAX_SEC", "15 / 300", "Exponential backoff on throttling"],
        ["MAX_RETRIES", "4", "Retries per call on throttle / transient errors"],
        ["THROTTLE_GAP_STEP / MAX / DECAY_SEC", "4 / 30 / 1", "Adaptive report spacing"],
        ["SWEEP_PASSES", "3", "Passes per date; later passes retry failures only"],
        ["PROGRESS_REFRESH_AFTER_HOURS", "4", "Re-sweep a finished commodity after this long"],
        ["STALE_EMPTY_WINDOW_DAYS / RECHECK_DAYS", "10 / 7", "Dormant-commodity skip"],
        ["MIRROR_PRUNE_STALE / MAX_PRUNE_RATIO", "true / 0.34", "Delete rows the site dropped; safety brake"],
        ["BLOB_PREFIX / AZURE_STORAGE_CONTAINER", "agmarknet_api/ / (VM)", "Where CSVs are mirrored"],
    ], widths=[2.4, 1.3, 2.9])

    # ---- 5. operations
    add_heading(doc, "5. Everyday commands", 1)
    add_code_block(doc, """python -m mandi_pipeline.run --once                       # one sweep of today
python -m mandi_pipeline.run --once --catchup-days 3      # today + previous 3 dates
python -m mandi_pipeline.run --status                     # last runs + today's progress
python -m mandi_pipeline.run --refresh-taxonomy           # re-pull the commodity list
python -m mandi_pipeline.run --verify-csv --date 2026-09-06 --deep --fix
python -m mandi_pipeline.run --verify-csv --date 2026-09-06 --blob
python -m mandi_pipeline.run --ship-logs
journalctl -u mandi-pipeline -f                           # live sweep log on the VM""")

    # ---- 6. flow
    add_heading(doc, "6. End-to-end flow", 1)
    add_body(doc, "Timers start a sweep → lock and run record → taxonomy, dormant skip, dates, resume → for each "
                  "commodity: CAPTCHA, report, parse, store, progress, CSV, blob (with retries for rejected "
                  "CAPTCHAs and throttling) → run summary, heartbeat or alert. The stored data is then checked "
                  "against the site by verify-csv and shown on the QA dashboard.")
    doc.add_picture(str(DIAGRAM_PATH), width=Inches(6.4))

    doc.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
