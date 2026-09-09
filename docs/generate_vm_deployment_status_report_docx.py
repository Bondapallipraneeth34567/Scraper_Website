"""Renders the Approach 1 VM deployment + connectivity incident report:
what was deployed and verified on sethu-admin1, the manual-fallback CAPTCHA
backup added afterward, the outage timeline, and what's pending.

Run:  python docs/generate_vm_deployment_status_report_docx.py
Out:  docs/vm_deployment_status_report.docx
"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_PATH = Path(__file__).parent / "vm_deployment_status_report.docx"

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x5C, 0x3D)
DARK = RGBColor(0x22, 0x22, 0x22)
WARN = RGBColor(0xA0, 0x3D, 0x1F)
GOOD = RGBColor(0x1F, 0x5C, 0x3D)
MUTED = RGBColor(0x55, 0x55, 0x55)


def _shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill}))


def code(doc, text):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        _shade(p, "F2F2F2")
        r = p.add_run(line if line.strip() else " ")
        r.font.name = CODE_FONT
        r.font.size = Pt(9)
        r.font.color.rgb = DARK


def body(doc, text, size=10.5, italic=False, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = BODY_FONT
    r.font.size = Pt(size)
    r.italic = italic
    return p


def bullet(doc, label, rest=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if label:
        r = p.add_run(label)
        r.bold = True
        r.font.name = BODY_FONT
        r.font.size = Pt(10.5)
        r.font.color.rgb = ACCENT
    r2 = p.add_run(rest)
    r2.font.name = BODY_FONT
    r2.font.size = Pt(10.5)


def num(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = BODY_FONT
    r.font.size = Pt(10.5)


def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = ACCENT
        r.font.name = BODY_FONT
    return h


def callout(doc, prefix, text, color):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    _shade(p, "F4F6F4")
    r1 = p.add_run(prefix + "  ")
    r1.bold = True
    r1.font.color.rgb = color
    r1.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.name = BODY_FONT
    r2.font.size = Pt(10.5)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = BODY_FONT
        r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.name = BODY_FONT
            r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def build():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    for s in doc.sections:
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    t = doc.add_heading("Approach 1 — VM Deployment & Connectivity Incident", level=0)
    for r in t.runs:
        r.font.color.rgb = ACCENT
    sub = doc.add_paragraph()
    sr = sub.add_run(
        "Deploying the AGMARKNET direct-API pipeline to sethu-admin1 (104.211.91.182), adding a "
        "manual CAPTCHA fallback afterward, and the VM going unresponsive mid-deployment. "
        "Written 4 September 2026, while the VM is still down."
    )
    sr.italic = True
    sr.font.size = Pt(11)
    sr.font.color.rgb = MUTED

    # ---- bottom line
    heading(doc, "Bottom line", 1)
    body(doc,
         "Approach 1's code, database schema, and systemd services were successfully deployed to "
         "sethu-admin1 and verified working from the VM itself (same 778 rows, same run history as "
         "the local build). Partway through a bounded live smoke test, the VM stopped responding — "
         "first SSH hung, then timed out, then the TCP connection was actively reset, and now even a "
         "raw TCP connect to port 22 times out. This looks like a VM-level or network-level outage, "
         "not an application bug. The existing production service on the same box "
         "(agmarknet-scraper-v2, writing mandi.market_prices_v2) is also unreachable while this "
         "persists. A restart requires Azure Portal/CLI access, which this session does not have — "
         "action is needed from the team.")

    # ---- what was deployed
    heading(doc, "1. What was deployed (and verified working)", 1)
    bullet(doc, "Target: ", "sethu-admin1 (Azure VM, 104.211.91.182), user azureuser, "
           "reused deploy key sethu-admin1_key.pem.")
    bullet(doc, "Scope: ", "Approach 1 ONLY (mandi_pipeline/agmarknet/ + shared plumbing: config.py, "
           "db/, logging_setup.py, normalize_common.py, storage.py, run.py, operator_console.py, "
           "deploy/). Approach 2 (datagovin/) was deliberately NOT pushed.")
    bullet(doc, "Code path: ", "/opt/mandi_pipeline/mandi_pipeline/, with its own .env "
           "(real Postgres creds matching backend/.env, CAPTCHA_SOLVER=easyocr, "
           "PROGRESS_REFRESH_AFTER_HOURS=3, Azure Blob creds pulled from the existing "
           "~/Web_Scraper/.env so the CSV mirror works).")
    bullet(doc, "Python deps: ", "installed into the EXISTING shared venv (~/Web_Scraper/venv) "
           "rather than a new dedicated one — the VM disk was already 81% full (5.5 GB free) and "
           "that venv already carried easyocr/torch/opencv/numpy/loguru/python-dotenv/"
           "azure-storage-blob (5.3 GB). Only httpx and psycopg (a few MB) were added.")
    bullet(doc, "Database: ", "python -m mandi_pipeline.run --init-db applied from the VM confirmed "
           "reachable and idempotent — mandi.pipeline_runs / pipeline_progress / agmarknet_taxonomy / "
           "market_prices_v3 already existed from the local build and were left as-is.")
    bullet(doc, "systemd: ", "mandi-pipeline.service/.timer and mandi-pipeline-console.service "
           "installed, daemon-reloaded, and enabled --now. ExecStart points at "
           "~/Web_Scraper/venv/bin/python (documented in the unit file) to match the shared-venv "
           "decision above.")
    bullet(doc, "Verification from the VM itself: ", "run --status returned the SAME run history and "
           "row count (778 rows for 2026-09-04) the local machine had built — confirms the VM reaches "
           "the same shared Postgres correctly.")
    body(doc, "Confirmed before the outage:", after=4)
    table(doc, ["Check", "Result"], [
        ["Timer enabled/active", "mandi-pipeline.timer: enabled, next fire 16:07 UTC"],
        ["Console enabled/active", "mandi-pipeline-console.service: enabled, active (running)"],
        ["DB reachable from VM", "yes -- same 778 rows / same 4 prior runs visible"],
        ["Deps importable", "httpx, psycopg, loguru, dotenv, easyocr, cv2 all import cleanly"],
    ], widths=[2.4, 4.0])

    # ---- captcha backup
    heading(doc, "2. CAPTCHA backup: audio investigated, manual fallback built instead", 1)
    body(doc, "Asked for an audio CAPTCHA as a backup path. Verified live against the real API "
              "rather than assuming:")
    code(doc,
         "POST /captcha/generator {type:'audio'}    -> HTTP 200, still returns a PNG (param ignored)\n"
         "POST /captcha/generator {format:'audio'}  -> HTTP 200, still returns a PNG (param ignored)\n"
         "GET  /captcha/audio                       -> HTTP 404\n"
         "POST /captcha/audio                       -> HTTP 404\n"
         "GET  /captcha/generator/audio              -> HTTP 404\n"
         "live form page HTML                        -> zero mentions of \"audio\"")
    callout(doc, "Finding:", "AGMARKNET's CAPTCHA is image-only. There is no audio alternative to "
            "fall back to.", WARN)
    body(doc, "Built the real equivalent instead: when the configured automated solver "
              "(easyocr / llm / llm_assisted_ocr) exhausts its 4 attempts on a commodity, the "
              "pipeline now escalates to ONE manual (human, operator-console) attempt before giving "
              "up, instead of marking the commodity failed outright. New setting "
              "CAPTCHA_MANUAL_FALLBACK (default true); no-op if CAPTCHA_SOLVER=manual already.")
    body(doc, "Proven live (forced the primary solver to fail against the real API):")
    code(doc,
         "AlwaysFailsSolver -> 4 real CAPTCHA images minted and rejected\n"
         "-> \"automated CAPTCHA solving exhausted -> escalating to manual fallback "
         "(there is no audio alternative -- verified, AGMARKNET's CAPTCHA is image-only)\"\n"
         "-> pending image written to state/captcha/, console waits\n"
         "-> no operator answered within the test timeout -> recorded 'failed' with that exact reason")
    callout(doc, "Status:", "Built and verified locally against the live API. NOT yet on the VM -- "
            "the outage started right after this was finished, before it could be pushed.", WARN)

    # ---- outage timeline
    doc.add_page_break()
    heading(doc, "3. Outage timeline", 1)
    table(doc, ["Time (UTC, approx)", "Event"], [
        ["13:34:41", "run10 started on the VM: bounded smoke test (--only Onion, trigger=manual)."],
        ["~13:35-13:44", "SSH commands to the VM start taking >120s to return (previously instant)."],
        ["13:44", "Direct DB query shows run10 still status=running, ~10 minutes in -- "
         "a 1-commodity run normally finishes in well under a minute."],
        ["shortly after", "SSH attempt 1: \"Connection timed out during banner exchange.\""],
        ["", "SSH attempt 2: same timeout."],
        ["", "SSH attempt 3: \"kex_exchange_identification: read: Connection reset by peer\" "
         "-- a DIFFERENT failure mode (TCP connects, then the handshake is reset)."],
        ["", "SSH attempts since: back to \"Connection timed out during banner exchange.\""],
        ["most recent check", "Raw TCP connect to port 22 (bypassing the SSH client entirely) also "
         "times out -- no response at the network level at all."],
    ], widths=[1.6, 5.2])
    callout(doc, "Not yet confirmed (no portal/serial-console access from this session):",
            "the root cause. The leading hypothesis is memory pressure -- the smoke test was the "
            "first time this session loaded easyocr's model (~torch) on the VM, on a box already "
            "running agmarknet-scraper-v2.service and mandi-voice-bot.service with only 5.5 GB disk "
            "free -- but this is a guess, not a diagnosis. The changing failure signature (timeout -> "
            "reset -> timeout -> no TCP response) is also consistent with a reboot in progress, a "
            "security group / firewall change, or an Azure-side network event.", WARN)

    # ---- what this affects
    heading(doc, "4. What this affects", 1)
    bullet(doc, "agmarknet-scraper-v2.service ", "(the existing production browser scraper, writing "
           "mandi.market_prices_v2) is on the SAME VM and is equally unreachable right now -- not "
           "just mandi_pipeline.")
    bullet(doc, "mandi-voice-bot.service ", "(the Telugu voice agent) is also on this VM.")
    bullet(doc, "mandi-pipeline.timer's first real 4-hourly cycle ", "(due 16:07 UTC) will not run "
           "while the VM is down; Persistent=true means it fires as soon as the VM is back, "
           "not that it silently skips.")
    bullet(doc, "The manual-fallback CAPTCHA code ", "and a couple of small config touches are only "
           "on the local machine, not the VM.")
    bullet(doc, "run10 ", "will need its pipeline_runs row cleaned up (marked aborted) once reachable "
           "again -- it cannot have finished on its own.")

    # ---- what's needed
    heading(doc, "5. What's needed from the team", 1)
    num(doc, "Restart/investigate the VM via Azure Portal (Virtual machines -> sethu-admin1 -> "
             "Restart) or Azure CLI -- this session has neither.")
    num(doc, "If it doesn't come back cleanly, use the Portal's Serial console (Support + "
             "troubleshooting) to get a shell without relying on SSH/networking, and check "
             "dmesg / journalctl for an OOM kill or other crash signal around 13:34-13:44 UTC "
             "4 Sep 2026.")
    num(doc, "Once reachable again: confirm agmarknet-scraper-v2.service and mandi-voice-bot.service "
             "came back healthy (they matter more than this pipeline).")
    num(doc, "Then: push the manual-fallback CAPTCHA code, clean up run10, and re-verify the timer "
             "and console service are enabled and active before leaving it unattended again.")

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
